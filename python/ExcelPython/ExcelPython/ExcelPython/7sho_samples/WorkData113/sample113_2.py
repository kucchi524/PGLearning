import openpyxl
import docx
fname = '販売記録08.xlsx'
wb = openpyxl.load_workbook(fname)
ws = wb.active
doc = docx.Document('販売報告01.docx')
tbl = doc.tables[0]
for i, row in enumerate(tbl.rows):
    if i > 0:
        line = []
        for cell in row.cells:
            val = cell.text
            cval = int(val) if val.isdigit() else val
            line.append(cval)
        ws.append(line)
wb.save(fname)
