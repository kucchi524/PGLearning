import openpyxl
fname = '注文記録10.xlsx'
wb = openpyxl.load_workbook(fname)
ws = wb.active
ws.tables['注文'].ref = 'A3:E10'
ws.tables['注文'].displayName = '記録'
wb.save(fname)
